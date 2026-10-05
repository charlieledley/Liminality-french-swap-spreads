# -*- coding: utf-8 -*-
"""Ingest Charlie's IRR / MOIC scenario output and tie it to the spread history.

Source: analysis/model/french_ss_details_and_returns_2026-10-05.xlsx (Charlie's spreadsheet
output, received 2026-10-05, unchanged). The IRR and MOIC are computed there, not here. This
script reads them, checks they are mutually consistent, places each scenario on the
percentile ladder from spread_history.py, and writes the deck data and an audit workbook.

What the spreadsheet gives (Sheet1):
  row 2  at-expiry 1y spread (bp, vs 6m Euribor swap, swap minus bond)
  row 3  IRR, 4-year, gross of fees            (blank where MOIC < 1)
  row 4  MOIC, multiple of invested capital
  B7:C21 model inputs: trade date, option expiry, funding spread, bond, entry spread, floor,
         fee basis, capital per 100m notional

Convention found in the data: IRR = MOIC ** (1 / 4.420) - 1 at every point (implied horizon
4.420 years against a stated 4.403 from trade date to expiry). Where the spreadsheet leaves
IRR blank (MOIC below 1) this script derives it with the same convention and marks it so.

Charlie, 2026-10-05: the roll-to-spot scenario uses a 1y spread of -1.3bp (interpolated to
the bond's maturity at expiry), not the +6bp 1y CMT point. S4 is read by linear interpolation
between the 0 and -25 grid points.
"""
import json
import os
import sys

import numpy as np
import openpyxl
import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(os.path.expanduser("~"), ".claude", "skills", "xlsx-report"))
from xlsx_report import Block, write_report          # noqa: E402

ASOF = "2026-10-05"
MODEL = os.path.join(HERE, "model", "french_ss_details_and_returns_2026-10-05.xlsx")
HISTORY = os.path.join(HERE, "..", "deck", "data", "spread_history_2026-10-05.json")
POINTS = os.path.join(HERE, "..", "deck", "data", "scenario_points_2026-10-05.json")
ROLL_TO_SPOT_BP = -1.3    # Charlie, 2026-10-05


def read_model():
    ws = openpyxl.load_workbook(MODEL, data_only=True)["Sheet1"]
    grid = []
    for c in range(5, ws.max_column + 1):
        s = ws.cell(2, c).value
        if s is None:
            continue
        grid.append({"spread_bp": float(s), "irr": ws.cell(3, c).value, "moic": float(ws.cell(4, c).value)})
    inputs = {}
    for r in range(7, 22):
        k, v = ws.cell(r, 2).value, ws.cell(r, 3).value
        if k is None:
            continue
        if hasattr(v, "date"):
            v = str(v.date())
        inputs[k] = v
    return grid, inputs


def check_convention(grid):
    """IRR and MOIC should be one horizon apart: ln(MOIC)/ln(1+IRR) constant."""
    T = [np.log(g["moic"]) / np.log(1 + g["irr"]) for g in grid if g["irr"] and g["irr"] > 0 and g["moic"] > 0]
    T = np.array(T)
    assert T.std() < 1e-3, f"IRR/MOIC horizon is not constant: {T}"
    return float(T.mean())


def interp(grid, x, key):
    """Linear between grid points; below the full-impairment level the return is a total loss, not extrapolated."""
    xs = np.array([g["spread_bp"] for g in grid])
    ys = np.array([g[key] if g[key] is not None else np.nan for g in grid], dtype=float)
    o = np.argsort(xs)
    if x < xs.min():
        return 0.0 if key == "moic" else -1.0
    return float(np.interp(x, xs[o], ys[o]))


def build():
    grid, inputs = read_model()
    T = check_convention(grid)
    for g in grid:
        if g["irr"] is None:
            g["irr"] = float(g["moic"] ** (1 / T) - 1) if g["moic"] > 0 else -1.0
            g["irr_source"] = "derived: MOIC^(1/T)-1, same convention as the spreadsheet"
        else:
            g["irr"] = float(g["irr"])
            g["irr_source"] = "spreadsheet"
    hist = json.load(open(HISTORY, encoding="utf-8"))
    pts = json.load(open(POINTS, encoding="utf-8"))
    lad = hist["by_basis"]["eur6m"]["ladder_1y"][hist["pct_start"][:4]]
    # percentile rank of a level on the since-2010 1y ladder, by interpolation on the ladder's own points
    pk = sorted(int(k[1:]) for k in lad if k.startswith("p"))
    lx = [lad["min"]] + [lad[f"p{p}"] for p in pk] + [lad["max"]]
    ly = [0] + pk + [100]

    def pct(level):
        return float(np.interp(level, lx, ly)) if lx[0] <= level <= lx[-1] else (0.0 if level < lx[0] else 100.0)

    on_grid = {g["spread_bp"] for g in grid}
    italy = hist["by_basis"]["eur6m"]["italy"]["1y"]
    named = []
    for n in pts["named"]:
        s = float(n["spread_bp"])
        n = dict(n)
        if n["id"] == "BE":
            s = next(g["spread_bp"] for g in grid if g["moic"] == 1.0)
        if n["id"] == "IMP":
            s = next(g["spread_bp"] for g in grid if g["moic"] == 0.0)
        if n["id"] == "S1":      # median of the 1y ladder on the decided construction, rounded to the bp
            s = float(round(lad["p50"]))
            n["from"] = "1y ladder p50 (%.1f)" % lad["p50"]
        if n["id"] == "S2":
            s = float(round(lad["p10"]))
            n["from"] = "1y ladder p10 (%.1f)" % lad["p10"]
        if n["id"] == "S5":      # Italy's stress level: the day count comes from the history, not the JSON
            s = float(italy["stress_level_bp"])
            a, b = italy["days_below_stress_span"]
            n["label"] = ("Italy's 2011 crisis level: the 1y point was below %dbp on only %d trading days, %s to %s"
                          % (s, italy["days_below_stress"], a, b))
        if n["id"] == "S6":      # Italy's euro-crisis low, read from the history on the decided construction
            s = round(float(italy["min"]))
            n["label"] = "Italy's 1y point at its euro-crisis low, %s" % italy["min_date"]
            n["from"] = "Italy 1y min 2010-13 (%.1f, %s)" % (italy["min"], italy["min_date"])
        named.append({"id": n["id"], "label": n["label"], "spread_bp": s,
                      "irr": interp(grid, s, "irr"), "moic": interp(grid, s, "moic"),
                      "pct_since_2010": pct(s), "from": n["from"],
                      "return_source": "spreadsheet grid" if s in on_grid else "linear interpolation between grid points"})
    named.sort(key=lambda n: -n["spread_bp"])      # richest first; the Italy low can sit below full impairment
    # the executive-summary claim: how often Italy and Spain breached break-even, on the deck's basis
    be = next(n["spread_bp"] for n in named if n["id"] == "BE")
    sys.path.insert(0, HERE)
    import load_rates
    w = load_rates.load()
    breach = {}
    for c in ("ita", "esp"):
        s = w[f"{c}_eur6m_1y_bp"].dropna()
        bl = s[s < be]
        breach[c] = {"first": str(s.index[0].date()), "last": str(s.index[-1].date()), "min": float(s.min()),
                     "min_date": str(s.idxmin().date()), "days_below_breakeven": int(len(bl)),
                     "span": ([str(bl.index.min().date()), str(bl.index.max().date())] if len(bl) else None)}
    out = {"received": ASOF, "model_file": os.path.relpath(MODEL, os.path.join(HERE, "..")),
           "basis": "6m Euribor swap, swap minus bond, bp", "fee_basis": inputs.get("Gross of Fees", "Gross of Fees"),
           "irr_moic_convention": f"IRR = MOIC^(1/{T:.3f}) - 1 at every spreadsheet point (implied horizon {T:.3f}y; stated {inputs['Yrs Time to Expiry']:.3f}y)",
           "model_inputs": inputs, "roll_to_spot_bp": ROLL_TO_SPOT_BP,
           "grid": sorted(grid, key=lambda g: -g["spread_bp"]), "named": named,
           "percentile_window_start": hist["pct_start"], "breakeven_breaches_1y": breach}
    return out


def workbook(out, path):
    g = pd.DataFrame(out["grid"])
    g["IRR (%)"] = g["irr"] * 100
    grid_df = g.rename(columns={"spread_bp": "1y spread at expiry (bp)", "moic": "MOIC", "irr_source": "IRR source"})[
        ["1y spread at expiry (bp)", "IRR (%)", "MOIC", "IRR source"]]
    n = pd.DataFrame(out["named"])
    n["IRR (%)"] = n["irr"] * 100
    named_df = n.rename(columns={"id": "#", "label": "Scenario", "spread_bp": "1y spread at expiry (bp)", "moic": "MOIC",
                                 "pct_since_2010": "Percentile since 2010 (%)", "from": "Level from",
                                 "return_source": "Return from"})[
        ["#", "Scenario", "1y spread at expiry (bp)", "IRR (%)", "MOIC", "Percentile since 2010 (%)", "Level from", "Return from"]]
    inp = pd.DataFrame({"Input": list(out["model_inputs"]), "Value": [str(v) for v in out["model_inputs"].values()]})
    sheets = {
        "Scenarios": [
            Block("Named scenarios: 4y IRR and MOIC by at-expiry 1y spread (vs 6m Euribor, swap minus bond)", named_df,
                  label_col="#", formats={"1y spread at expiry (bp)": "num1", "IRR (%)": "pct1", "MOIC": "num2",
                                          "Percentile since 2010 (%)": "num0", "Scenario": "text", "Level from": "text",
                                          "Return from": "text"},
                  note="IRR and MOIC from Charlie's spreadsheet (gross of fees), linearly interpolated between grid points where a "
                       "scenario falls off the grid. Percentile = position on the since-2010 daily ladder of the 1y France spread, low = cheap."),
            Block("Model inputs, as used in the spreadsheet", inp, label_col="Input", formats={"Value": "text"},
                  note=out["irr_moic_convention"]),
        ],
        "Grid": [Block("IRR and MOIC grid", grid_df, label_col="1y spread at expiry (bp)",
                       formats={"IRR (%)": "pct1", "MOIC": "num2", "IRR source": "text"},
                       chart={"kind": "line", "values": ["MOIC"], "title": "MOIC by at-expiry 1y spread"})],
    }
    write_report(path, sheets, title="French swap spread: scenario returns",
                 subtitle=f"Spreadsheet output received {ASOF}; analysis/scenario_returns.py")


if __name__ == "__main__":
    out = build()
    jp = os.path.join(HERE, "..", "deck", "data", f"scenario_irr_moic_{ASOF}.json")
    with open(jp, "w", encoding="utf-8") as f:
        json.dump(out, f, indent=1, ensure_ascii=False)
    xp = os.path.join(HERE, "..", "exports", f"scenario_irr_moic_{ASOF}.xlsx")
    workbook(out, xp)
    print("wrote", os.path.relpath(jp, os.path.join(HERE, "..")), "and", os.path.relpath(xp, os.path.join(HERE, "..")))
    print(out["irr_moic_convention"])
    print(f"\n{'#':4s} {'scenario':48s} {'spread':>7s} {'IRR':>7s} {'MOIC':>6s} {'pct':>5s}")
    for n in out["named"]:
        print(f"{n['id']:4s} {n['label'][:48]:48s} {n['spread_bp']:7.1f} {n['irr']*100:6.1f}% {n['moic']:6.2f} {n['pct_since_2010']:5.1f}")
